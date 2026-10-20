import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from jobs import swing_performance
from scanner.signal_history import IST, SignalHistory


class SignalHistoryTests(unittest.TestCase):
    def test_legacy_swing_counts_migrate_without_claiming_trading_day_history(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "swing.csv"
            path.write_text(
                "Alert Type,Ticker Name,Date Time,Buy Price,Signals On Entry Day,Signals In Last 7 Days\n"
                "EP,TEST,2026-10-19T09:16:00+05:30,100,4,6\n",
                encoding="utf-8",
            )
            with patch.object(swing_performance, "OUTPUT_FILE", path):
                row = next(iter(swing_performance._read_existing().values()))
            self.assertEqual(row["Intraday Alert Count"], "4")
            self.assertEqual(row["Daily Alert Count (7 Trading Days)"], "NA")
            self.assertNotIn("Signals In Last 7 Days", row)

    def test_rolling_counts_all_types_restart_dedup_and_ist_boundaries(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "history.csv"
            sessions = path.with_name("trading_sessions.csv")
            history = SignalHistory(path, sessions)
            history.record_sessions([
                f"{day}T09:15:00+05:30"
                for day in ("2026-10-09", "2026-10-12", "2026-10-13", "2026-10-14",
                            "2026-10-15", "2026-10-16", "2026-10-19", "2026-10-20")
            ])
            for timestamp, kind in (
                ("2026-10-09T09:16:00+05:30", "EP"),
                ("2026-10-12T09:16:00+05:30", "EP"),
                ("2026-10-13T09:16:00+05:30", "EP"),
                ("2026-10-19T09:16:00+05:30", "EP"),
                ("2026-10-19T09:17:00+05:30", "BUY"),
            ):
                message = history.record("test", kind, datetime.fromisoformat(timestamp))
            self.assertIn("Intraday Alert Count: 2 today", message)
            self.assertIn("Daily Alert Count: 4 of last 7 trading days", message)
            restarted = SignalHistory(path, sessions)
            self.assertEqual(restarted.counts("TEST", datetime(2026, 10, 19).date()), (2, 4))
            # Importing the same artifact again does not count events twice.
            restarted.extend(history.events)
            self.assertEqual(len(restarted.events), 5)
            message = restarted.record(
                "TEST", "EP_ORB_30M", datetime.fromisoformat("2026-10-19T09:46:00+05:30"),
            )
            self.assertIn("Intraday Alert Count: 3 today", message)
            self.assertIn("Daily Alert Count: 4 of last 7 trading days", message)
            restarted.record("OTHER", "BUY", datetime(2026, 10, 19, 10, tzinfo=IST))
            restarted.record("TEST", "EP", datetime.fromisoformat("2026-10-19T20:00:00+00:00"))
            self.assertEqual(restarted.counts("TEST", datetime(2026, 10, 19).date()), (3, 4))
            self.assertEqual(restarted.counts("TEST", datetime(2026, 10, 20).date()), (1, 4))

    def test_daily_alert_count_is_na_without_trading_session_calendar(self):
        with TemporaryDirectory() as directory:
            history = SignalHistory(Path(directory) / "history.csv")
            message = history.record("TEST", "EP", datetime(2026, 10, 20, 9, 16, tzinfo=IST))
            self.assertIn("Intraday Alert Count: 1 today", message)
            self.assertIn("Daily Alert Count: NA", message)

    def test_invalid_history_is_reported_not_silently_reset(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "history.csv"
            path.write_text("symbol,alert_time_ist,alert_type\nTEST,bad,EP\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Invalid signal-history timestamp"):
                SignalHistory(path)

    def test_swing_preserves_first_entry_sector_and_counts_all_daily_signals(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            history_path = root / "history.csv"
            history = SignalHistory(history_path, root / "trading_sessions.csv")
            history.record_sessions([
                f"{day}T09:15:00+05:30"
                for day in ("2026-10-09", "2026-10-12", "2026-10-13", "2026-10-14",
                            "2026-10-15", "2026-10-16", "2026-10-19")
            ])
            history.record("TEST", "EP", datetime(2026, 10, 16, 9, 16, tzinfo=IST))
            alerts = [{
                "Alert ID": kind, "Ticker Name": "TEST", "Buy Price": price,
                "Date Time": f"2026-10-19T{time}+05:30", "Alert Type": kind,
                "Sector": "Refineries", "Market Context": "Bull Market",
            } for kind, price, time in (
                ("EP", "100", "09:16:00"), ("BUY", "110", "09:20:00"),
                ("EP_ORB_30M", "112", "09:46:00"),
            )]
            with (
                patch.object(swing_performance, "SIGNAL_HISTORY_FILE", history_path),
                patch.object(swing_performance, "_read_existing", return_value={}),
                patch.object(swing_performance, "_read_current_alerts", return_value=alerts),
                patch.object(swing_performance, "ARTIFACT_DIR", root),
                patch.object(swing_performance, "OUTPUT_FILE", root / "swing.csv"),
                patch.object(swing_performance, "load_nse_equities", return_value=pd.DataFrame({
                    "trading_symbol": ["TEST"], "instrument_key": ["KEY"],
                })),
                patch.object(swing_performance, "_evaluate_alert",
                             side_effect=lambda row, key: {**row, "Status": "Trade open"}),
                patch.object(swing_performance, "_write_excel") as excel,
                patch.object(swing_performance, "_send_exit_alerts"),
                patch.object(swing_performance, "_send_excel_report"),
            ):
                swing_performance.run()
            frame = excel.call_args.args[0]
            self.assertEqual(len(frame), 1)
            row = frame.iloc[0]
            self.assertEqual(row["Buy Price"], "100")
            self.assertEqual(row["Sector"], "Refineries")
            self.assertEqual(row["Intraday Alert Count"], 3)
            self.assertEqual(row["Daily Alert Count (7 Trading Days)"], 2)


class RegimeDashboardTests(unittest.TestCase):
    def write_and_read(self, rows, root):
        frame = pd.DataFrame(rows, columns=swing_performance.OUTPUT_COLUMNS)
        path = root / "swing.xlsx"
        with patch.object(swing_performance, "EXCEL_FILE", path):
            swing_performance._write_excel(frame)
        return load_workbook(path)

    def test_regimes_use_closed_only_win_rate_and_keep_open_returns_separate(self):
        rows = []
        for index, (regime, status, value) in enumerate((
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Full exit", "10"),
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Full exit", "-5"),
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Full exit", "0"),
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Trade open", "20"),
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Partial exit", "5"),
            ("Market Context: Bear Market\nNIFTY 500: Downtrending\nStocks above 200 MA: 40%", "Full exit", "-10"),
            ("Market Context: Mixed Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 40%", "Trade open", "2"),
            ("", "ERROR", ""),
            ("Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60%", "Full exit", ""),
        )):
            row = dict.fromkeys(swing_performance.OUTPUT_COLUMNS, "")
            row.update({
                "Ticker Name": f"TEST{index}", "Alert Date": "2026-10-18",
                "Alert Time": "09:16:00", "Market Context": regime,
                "Status": status, "Performance %": value, "Sector": "Refineries",
                "Intraday Alert Count": 2, "Daily Alert Count (7 Trading Days)": 3,
            })
            rows.append(row)
        with TemporaryDirectory() as directory:
            workbook = self.write_and_read(rows, Path(directory))
            sheet = workbook["Dashboard"]
            self.assertEqual([sheet.cell(17, col).value for col in range(1, 13)], [
                "Bull Market", 6, 1, 1, 4, 1, 1, 1, "33.33%", "+6.00%", 0, 3,
            ])
            self.assertEqual(sheet["I18"].value, "0.00%")
            self.assertEqual(sheet["I19"].value, "NA")
            self.assertEqual(sheet["A20"].value, "Mixed Market")
            self.assertEqual(sheet["B20"].value, 1)
            self.assertEqual(sheet["A21"].value, "NA")
            self.assertEqual(sheet["A22"].value, "Unknown")
            self.assertEqual(sheet["K22"].value, 1)
            detail = workbook["Swing Performance"]
            self.assertEqual(detail["P3"].value, "Sector")
            self.assertEqual(detail["Q3"].value, "Intraday Alert Count")
            self.assertEqual(detail["R3"].value, "Daily Alert Count (7 Trading Days)")
            workbook.close()

    def test_empty_dashboard_has_all_regimes_and_no_fake_win_rate(self):
        with TemporaryDirectory() as directory:
            workbook = self.write_and_read([], Path(directory))
            sheet = workbook["Dashboard"]
            for row in range(17, 23):
                self.assertEqual(sheet.cell(row, 2).value, 0)
                self.assertEqual(sheet.cell(row, 9).value, "NA")
                self.assertEqual(sheet.cell(row, 10).value, "NA")
            workbook.close()

    def test_mixed_context_uses_non_bull_swing_profit_rules(self):
        self.assertEqual(swing_performance._market_rules(
            "Market Context: Mixed Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 40% (Broader Bear)"
        ), (10.0, 0.30))
        self.assertEqual(swing_performance._market_rules(
            "Market Context: Bull Market\nNIFTY 500: Uptrending\nStocks above 200 MA: 60% (Broader Bull)"
        ), (15.0, 0.20))


if __name__ == "__main__":
    unittest.main()
